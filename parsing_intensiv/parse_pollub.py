import re

from bs4 import BeautifulSoup
import requests
from openpyxl import Workbook
import time

from openpyxl.worksheet.table import Table, TableStyleInfo

base_url = "https://cs.pollub.pl/staff/"
def get_values(link):
    page = requests.get(link)
    soup_new = BeautifulSoup(page.content, "lxml")
    name = soup_new.select("article:first-child")[0].text
    article = soup_new.find("article")
    contact = article.find(lambda tag: tag.name == "p" and "e-mail" in tag.text)
    print(article)
    print(content)
    return [name, contact]

if __name__ == "__main__":
    dest_filaname = "parcownicy.xlsx"
    wb = Workbook()
    ws = wb.create_sheet(title="Pracownicy")
    ws.append(["Imię", "Nazwisko", "Kontakt"])
    staff = requests.get(base_url)
    soup = BeautifulSoup(staff.content, "lxml")
    content = soup.find("div", class_="post-content").find("article")
    links = content.find_all("a")
    for i in range(5):
        href = links[i]["href"]
        row = get_values(href)
        print(row)

        # ws.append(row)

    # tab = Table(displayName="Table1", ref="A1:E5")
    #
    # # Add a default style with striped rows and banded columns
    # style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
    #                        showLastColumn=False, showRowStripes=True, showColumnStripes=True)
    # tab.tableStyleInfo = style
    # ws.add_table(tab)
    # wb.save(dest_filaname)
    # print('Stuff content: {}'.format(content))
    # print('Stuff length: {}'.format(len(links)))
